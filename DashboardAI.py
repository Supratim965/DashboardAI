import html
from pathlib import Path
import warnings

import pandas as pd
import torch
from openpyxl import load_workbook
from transformers import AutoModelForCausalLM, AutoTokenizer, GenerationConfig

MODEL_ID = "TinyLlama/TinyLlama-1.1B-Chat-v1.0"
LOCAL_MODEL_DIR = Path(".hf-models") / "TinyLlama-1.1B-Chat-v1.0"
INPUT_FILE = Path("dashboard.xlsx")
OUTPUT_DATA_CSV = Path("ai_dashboard_data.csv")
OUTPUT_SUGGESTIONS_CSV = Path("ai_suggestions.csv")
OUTPUT_REGION_TOTALS_CSV = Path("region_totals.csv")
OUTPUT_HTML = Path("ai_dashboard.html")
OUTPUT_SITE_DIR = Path("docs")
OUTPUT_SITE_HTML = OUTPUT_SITE_DIR / "index.html"
OUTPUT_NOJEKYLL = OUTPUT_SITE_DIR / ".nojekyll"

DEFAULT_INSTRUCTIONS = [
    "1. Bar chart: Sales by Month (column A vs B)",
    "2. Pie chart: Sales by Region using total sales per region",
    "3. Summary table: Total sales by region",
]


def build_generator():
    """Load the local Hugging Face model when available."""
    model_dtype = torch.float16 if torch.cuda.is_available() else torch.float32
    model_kwargs = {"dtype": model_dtype}
    if torch.cuda.is_available():
        model_kwargs["device_map"] = "auto"

    model_source = str(LOCAL_MODEL_DIR) if LOCAL_MODEL_DIR.exists() else MODEL_ID

    try:
        tokenizer = AutoTokenizer.from_pretrained(model_source, local_files_only=True)
        model = AutoModelForCausalLM.from_pretrained(
            model_source,
            local_files_only=True,
            **model_kwargs,
        )
        if tokenizer.pad_token_id is None:
            tokenizer.pad_token = tokenizer.eos_token

        generation_config = GenerationConfig(
            max_new_tokens=120,
            max_length=None,
            do_sample=False,
            pad_token_id=tokenizer.pad_token_id,
            eos_token_id=tokenizer.eos_token_id,
        )
        return tokenizer, model, generation_config
    except Exception as exc:
        print(f"Warning: could not load Hugging Face model '{model_source}': {exc}")
        print("Falling back to built-in dashboard suggestions.")
        return None


def normalize_instructions(raw_text):
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    numbered_lines = [line for line in lines if len(line) > 2 and line[0].isdigit() and line[1] == "."]
    valid_lines = [
        line for line in numbered_lines
        if any(keyword in line.lower() for keyword in ("chart", "table", "summary"))
    ]
    return valid_lines[:3] if len(valid_lines) >= 3 else DEFAULT_INSTRUCTIONS


def generate_instructions(generator, data_summary):
    if generator is None:
        return DEFAULT_INSTRUCTIONS

    prompt = f"""Analyze this data summary:
{data_summary}

Return exactly 3 numbered dashboard suggestions.
Use this format only:
1. <chart or table suggestion>
2. <chart or table suggestion>
3. <chart or table suggestion>"""

    tokenizer, model, generation_config = generator
    inputs = tokenizer(prompt, return_tensors="pt")
    if torch.cuda.is_available():
        inputs = {key: value.to(model.device) for key, value in inputs.items()}

    prompt_token_count = inputs["input_ids"].shape[1]
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=r"Both `max_new_tokens`.*")
        outputs = model.generate(**inputs, generation_config=generation_config)
    generated_tokens = outputs[0][prompt_token_count:]
    response = tokenizer.decode(generated_tokens, skip_special_tokens=True).strip()
    return normalize_instructions(response)


def load_input_data():
    """Load source data from dashboard.xlsx if present, otherwise fall back to sample data."""
    if not INPUT_FILE.exists():
        return pd.DataFrame({
            "Month": ["Jan", "Feb", "Mar", "Apr"],
            "Sales": [100, 150, 120, 180],
            "Region": ["North", "South", "North", "South"],
        })

    workbook = load_workbook(INPUT_FILE, data_only=True)
    if "Data" not in workbook.sheetnames:
        raise ValueError(f"'{INPUT_FILE}' exists but does not contain a 'Data' sheet.")

    sheet = workbook["Data"]
    rows = list(sheet.iter_rows(values_only=True))
    data_rows = [list(row[:3]) for row in rows if any(value is not None for value in row[:3])]
    if not data_rows:
        raise ValueError(f"No usable rows were found in '{INPUT_FILE}' sheet 'Data'.")

    header = [str(value).strip() if value is not None else "" for value in data_rows[0]]
    expected = ["month", "sales", "region"]
    if [value.lower() for value in header] == expected:
        records = data_rows[1:]
    else:
        records = data_rows

    df = pd.DataFrame(records, columns=["Month", "Sales", "Region"])
    df = df.dropna(how="all")
    df["Month"] = df["Month"].astype(str).str.strip()
    df["Region"] = df["Region"].astype(str).str.strip()
    df["Sales"] = pd.to_numeric(df["Sales"], errors="coerce")
    df = df.dropna(subset=["Month", "Sales", "Region"]).reset_index(drop=True)

    if df.empty:
        raise ValueError(f"No valid Month/Sales/Region rows were found in '{INPUT_FILE}'.")
    return df


def format_sales(value):
    return f"{value:,.0f}"


def build_sales_chart(df):
    chart_width = 560
    chart_height = 280
    left = 56
    right = 18
    top = 20
    bottom = 42
    inner_width = chart_width - left - right
    inner_height = chart_height - top - bottom
    max_sales = max(float(df["Sales"].max()), 1.0)
    count = len(df)
    slot_width = inner_width / max(count, 1)
    bar_width = min(52, slot_width * 0.58)

    bars = []
    labels = []
    values = []
    grid_lines = []
    for step in range(5):
        grid_value = max_sales * step / 4
        y = top + inner_height - (grid_value / max_sales) * inner_height
        grid_lines.append(
            f'<line x1="{left}" y1="{y:.1f}" x2="{chart_width - right}" y2="{y:.1f}" />'
            f'<text x="{left - 10}" y="{y + 4:.1f}">{format_sales(grid_value)}</text>'
        )

    for index, row in enumerate(df.itertuples(index=False), start=0):
        sales = float(row.Sales)
        x = left + slot_width * index + (slot_width - bar_width) / 2
        bar_height = (sales / max_sales) * inner_height
        y = top + inner_height - bar_height
        label_x = x + bar_width / 2
        bars.append(
            f'<rect x="{x:.1f}" y="{y:.1f}" width="{bar_width:.1f}" height="{bar_height:.1f}" rx="14" />'
        )
        labels.append(
            f'<text x="{label_x:.1f}" y="{chart_height - 14}" class="axis-label">{html.escape(str(row.Month))}</text>'
        )
        values.append(
            f'<text x="{label_x:.1f}" y="{max(y - 10, 16):.1f}" class="value-label">{format_sales(sales)}</text>'
        )

    return f"""
    <svg viewBox="0 0 {chart_width} {chart_height}" class="sales-chart" role="img" aria-label="Sales by month">
      <defs>
        <linearGradient id="barGradient" x1="0" y1="0" x2="0" y2="1">
          <stop offset="0%" stop-color="#34d399" />
          <stop offset="100%" stop-color="#0f766e" />
        </linearGradient>
      </defs>
      <g class="grid">{''.join(grid_lines)}</g>
      <g class="bars">{''.join(bars)}</g>
      <g class="value-labels">{''.join(values)}</g>
      <g class="labels">{''.join(labels)}</g>
    </svg>
    """


def build_region_chart(region_totals):
    palette = ["#0f766e", "#f59e0b", "#dc2626", "#2563eb", "#7c3aed", "#16a34a"]
    total_sales = float(region_totals["Sales"].sum()) or 1.0
    cursor = 0.0
    segments = []
    legend_items = []
    for index, row in enumerate(region_totals.itertuples(index=False)):
        share = float(row.Sales) / total_sales
        next_cursor = cursor + share * 100
        color = palette[index % len(palette)]
        segments.append(f"{color} {cursor:.2f}% {next_cursor:.2f}%")
        legend_items.append(
            "<div class=\"legend-item\">"
            f"<span class=\"legend-swatch\" style=\"background:{color}\"></span>"
            f"<span>{html.escape(str(row.Region))}</span>"
            f"<strong>{share * 100:.1f}%</strong>"
            "</div>"
        )
        cursor = next_cursor

    chart_style = f"background: conic-gradient({', '.join(segments)});"
    return f"""
    <div class="donut-wrap">
      <div class="donut-chart" style="{chart_style}">
        <div class="donut-hole">
          <span>Total</span>
          <strong>{format_sales(total_sales)}</strong>
        </div>
      </div>
      <div class="legend">{''.join(legend_items)}</div>
    </div>
    """


def build_table(headers, rows):
    header_html = "".join(f"<th>{html.escape(str(header))}</th>" for header in headers)
    body_rows = []
    for row in rows:
        cells = "".join(f"<td>{html.escape(str(cell))}</td>" for cell in row)
        body_rows.append(f"<tr>{cells}</tr>")
    return (
        "<table><thead><tr>"
        + header_html
        + "</tr></thead><tbody>"
        + "".join(body_rows)
        + "</tbody></table>"
    )


def render_html_dashboard(df, region_totals, instructions):
    total_sales = float(df["Sales"].sum())
    average_sales = float(df["Sales"].mean())
    best_row = df.loc[df["Sales"].idxmax()]
    top_region = region_totals.sort_values("Sales", ascending=False).iloc[0]

    suggestion_items = "".join(
        f"<li>{html.escape(line)}</li>"
        for line in instructions
    )
    data_rows = [
        [row.Month, format_sales(float(row.Sales)), row.Region]
        for row in df.itertuples(index=False)
    ]
    region_rows = [
        [row.Region, format_sales(float(row.Sales))]
        for row in region_totals.itertuples(index=False)
    ]

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>AI Sales Dashboard</title>
  <style>
    :root {{
      --ink: #182126;
      --muted: #5d6a72;
      --paper: #f6f0e8;
      --panel: rgba(255, 252, 247, 0.82);
      --line: rgba(24, 33, 38, 0.12);
      --teal: #0f766e;
      --amber: #f59e0b;
      --coral: #dc2626;
      --sky: #2563eb;
      --shadow: 0 24px 60px rgba(24, 33, 38, 0.12);
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      font-family: "Trebuchet MS", "Segoe UI Variable Text", sans-serif;
      color: var(--ink);
      background:
        radial-gradient(circle at top left, rgba(15, 118, 110, 0.20), transparent 32%),
        radial-gradient(circle at top right, rgba(245, 158, 11, 0.22), transparent 26%),
        linear-gradient(160deg, #f9f3ea 0%, #f1efe8 55%, #e8f1ef 100%);
      min-height: 100vh;
    }}
    .shell {{
      max-width: 1180px;
      margin: 0 auto;
      padding: 36px 20px 48px;
    }}
    .hero {{
      padding: 32px;
      border-radius: 28px;
      background: linear-gradient(135deg, rgba(255,255,255,0.9), rgba(245, 248, 247, 0.72));
      border: 1px solid rgba(255,255,255,0.75);
      box-shadow: var(--shadow);
      position: relative;
      overflow: hidden;
    }}
    .hero::after {{
      content: "";
      position: absolute;
      inset: auto -90px -110px auto;
      width: 280px;
      height: 280px;
      border-radius: 999px;
      background: rgba(15, 118, 110, 0.10);
      filter: blur(4px);
    }}
    h1, h2 {{
      font-family: "Palatino Linotype", "Book Antiqua", Georgia, serif;
      margin: 0;
    }}
    h1 {{
      font-size: clamp(2rem, 4vw, 3.6rem);
      line-height: 0.95;
      max-width: 9ch;
    }}
    .hero p {{
      max-width: 58ch;
      color: var(--muted);
      font-size: 1rem;
      line-height: 1.6;
      margin: 16px 0 0;
    }}
    .kpis {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
      gap: 16px;
      margin-top: 24px;
    }}
    .card, .panel {{
      background: var(--panel);
      backdrop-filter: blur(14px);
      border: 1px solid rgba(255,255,255,0.8);
      border-radius: 24px;
      box-shadow: var(--shadow);
    }}
    .card {{
      padding: 18px 18px 16px;
    }}
    .eyebrow {{
      color: var(--muted);
      font-size: 0.78rem;
      letter-spacing: 0.12em;
      text-transform: uppercase;
    }}
    .metric {{
      margin-top: 8px;
      font-size: 1.9rem;
      font-weight: 700;
    }}
    .metric-note {{
      margin-top: 6px;
      color: var(--muted);
      font-size: 0.92rem;
    }}
    .grid {{
      display: grid;
      grid-template-columns: 1.35fr 1fr;
      gap: 18px;
      margin-top: 22px;
    }}
    .panel {{
      padding: 22px;
    }}
    .panel h2 {{
      font-size: 1.45rem;
      margin-bottom: 6px;
    }}
    .panel-copy {{
      color: var(--muted);
      margin: 0 0 14px;
    }}
    .sales-chart {{
      width: 100%;
      height: auto;
      display: block;
    }}
    .sales-chart .grid line {{
      stroke: rgba(24, 33, 38, 0.10);
      stroke-width: 1;
    }}
    .sales-chart .grid text {{
      fill: var(--muted);
      font-size: 11px;
      text-anchor: end;
    }}
    .sales-chart .bars rect {{
      fill: url(#barGradient);
    }}
    .sales-chart .axis-label {{
      fill: var(--muted);
      font-size: 12px;
      text-anchor: middle;
    }}
    .sales-chart .value-label {{
      fill: var(--ink);
      font-size: 11px;
      font-weight: 700;
      text-anchor: middle;
    }}
    .sales-chart defs linearGradient stop:first-child {{
      stop-color: #0f766e;
    }}
    .sales-chart defs linearGradient stop:last-child {{
      stop-color: #34d399;
    }}
    .donut-wrap {{
      display: grid;
      gap: 18px;
      align-items: center;
      justify-items: center;
      min-height: 100%;
    }}
    .donut-chart {{
      width: min(280px, 70vw);
      aspect-ratio: 1;
      border-radius: 50%;
      display: grid;
      place-items: center;
      box-shadow: inset 0 0 0 1px rgba(255,255,255,0.35);
    }}
    .donut-hole {{
      width: 56%;
      aspect-ratio: 1;
      border-radius: 50%;
      background: rgba(255, 251, 246, 0.92);
      display: grid;
      place-items: center;
      text-align: center;
      box-shadow: inset 0 0 0 1px rgba(24, 33, 38, 0.08);
    }}
    .donut-hole span {{
      color: var(--muted);
      font-size: 0.85rem;
      text-transform: uppercase;
      letter-spacing: 0.08em;
    }}
    .donut-hole strong {{
      font-size: 1.6rem;
      line-height: 1;
    }}
    .legend {{
      width: 100%;
      display: grid;
      gap: 10px;
    }}
    .legend-item {{
      display: grid;
      grid-template-columns: 14px 1fr auto;
      gap: 10px;
      align-items: center;
      padding: 10px 12px;
      border-radius: 14px;
      background: rgba(255,255,255,0.58);
      border: 1px solid rgba(24, 33, 38, 0.07);
    }}
    .legend-swatch {{
      width: 14px;
      height: 14px;
      border-radius: 999px;
    }}
    .suggestions {{
      display: grid;
      grid-template-columns: 0.9fr 1.1fr;
      gap: 18px;
      margin-top: 18px;
    }}
    ol {{
      margin: 10px 0 0;
      padding-left: 22px;
    }}
    li {{
      margin-bottom: 10px;
      line-height: 1.45;
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      margin-top: 10px;
      font-size: 0.95rem;
      overflow: hidden;
      border-radius: 16px;
    }}
    th, td {{
      padding: 12px 14px;
      border-bottom: 1px solid var(--line);
      text-align: left;
    }}
    th {{
      background: rgba(15, 118, 110, 0.09);
      font-size: 0.82rem;
      text-transform: uppercase;
      letter-spacing: 0.08em;
      color: var(--muted);
    }}
    tbody tr:last-child td {{
      border-bottom: 0;
    }}
    .footer-note {{
      margin-top: 18px;
      color: var(--muted);
      font-size: 0.92rem;
    }}
    @media (max-width: 900px) {{
      .grid, .suggestions {{
        grid-template-columns: 1fr;
      }}
      .hero {{
        padding: 24px;
      }}
    }}
  </style>
</head>
<body>
  <div class="shell">
    <section class="hero">
      <div class="eyebrow">Offline HTML Dashboard</div>
      <h1>Sales performance at a glance</h1>
      <p>This dashboard is generated directly from your local project data and CSV outputs. It is fully self-contained, so you can open it in any browser without Excel or internet access.</p>
      <div class="kpis">
        <div class="card">
          <div class="eyebrow">Total Sales</div>
          <div class="metric">{format_sales(total_sales)}</div>
          <div class="metric-note">Across {len(df)} reporting rows</div>
        </div>
        <div class="card">
          <div class="eyebrow">Average Sale</div>
          <div class="metric">{format_sales(average_sales)}</div>
          <div class="metric-note">Mean sales per month</div>
        </div>
        <div class="card">
          <div class="eyebrow">Best Month</div>
          <div class="metric">{html.escape(str(best_row["Month"]))}</div>
          <div class="metric-note">{format_sales(float(best_row["Sales"]))} sales</div>
        </div>
        <div class="card">
          <div class="eyebrow">Top Region</div>
          <div class="metric">{html.escape(str(top_region["Region"]))}</div>
          <div class="metric-note">{format_sales(float(top_region["Sales"]))} sales</div>
        </div>
      </div>
    </section>

    <section class="grid">
      <article class="panel">
        <h2>Sales by Month</h2>
        <p class="panel-copy">A direct month-to-month view of the source sales data.</p>
        {build_sales_chart(df)}
      </article>
      <article class="panel">
        <h2>Sales by Region</h2>
        <p class="panel-copy">Regional contribution split based on the aggregated sales totals.</p>
        {build_region_chart(region_totals)}
      </article>
    </section>

    <section class="suggestions">
      <article class="panel">
        <h2>AI Suggestions</h2>
        <p class="panel-copy">Structured dashboard recommendations generated from the local model.</p>
        <ol>{suggestion_items}</ol>
      </article>
      <article class="panel">
        <h2>Region Summary</h2>
        <p class="panel-copy">Chart-ready aggregated totals exported to <code>region_totals.csv</code>.</p>
        {build_table(["Region", "Sales"], region_rows)}
      </article>
    </section>

    <section class="panel" style="margin-top: 18px;">
      <h2>Source Data</h2>
      <p class="panel-copy">Cleaned records exported to <code>ai_dashboard_data.csv</code>.</p>
      {build_table(["Month", "Sales", "Region"], data_rows)}
      <div class="footer-note">Generated from local files in this project folder.</div>
    </section>
  </div>
</body>
</html>
"""


generator = build_generator()
df = load_input_data()
region_totals = df.groupby("Region", as_index=False)["Sales"].sum()

data_summary = (
    df.describe(include="all").fillna("").to_string()
    + "\nColumns: "
    + str(df.columns.tolist())
)
instructions = generate_instructions(generator, data_summary)

print("AI Suggestions:")
for line in instructions:
    print(line)

df.to_csv(OUTPUT_DATA_CSV, index=False)
pd.DataFrame({"Suggestion": instructions}).to_csv(OUTPUT_SUGGESTIONS_CSV, index=False)
region_totals.to_csv(OUTPUT_REGION_TOTALS_CSV, index=False)
html_output = render_html_dashboard(df, region_totals, instructions)
OUTPUT_HTML.write_text(html_output, encoding="utf-8")
OUTPUT_SITE_DIR.mkdir(exist_ok=True)
OUTPUT_SITE_HTML.write_text(html_output, encoding="utf-8")
OUTPUT_NOJEKYLL.write_text("", encoding="utf-8")

print(f"Data CSV saved: {OUTPUT_DATA_CSV.resolve()}")
print(f"Suggestions CSV saved: {OUTPUT_SUGGESTIONS_CSV.resolve()}")
print(f"Region totals CSV saved: {OUTPUT_REGION_TOTALS_CSV.resolve()}")
print(f"HTML dashboard saved: {OUTPUT_HTML.resolve()}")
print(f"Shareable site entry saved: {OUTPUT_SITE_HTML.resolve()}")
